from io import BytesIO

from docx import Document

from knowledge.reading.converter import convert_to_docling
from knowledge.reading.sections import split_into_sections

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

NOTES = b"""# Machine Learning

## Unit 3: Linear Regression

Simple linear regression fits a straight line.

| x | y |
| --- | --- |
| 1 | 3 |
| 2 | 5 |

## Unit 4: Logistic Regression

Despite its name, logistic regression is a classification algorithm.
"""


def test_sections_follow_headings_and_keep_tables_whole() -> None:
    sections = split_into_sections(convert_to_docling("text/markdown", NOTES))
    paths = [section.heading_path for section in sections]
    assert ["Machine Learning", "Unit 3: Linear Regression"] in paths
    assert ["Machine Learning", "Unit 4: Logistic Regression"] in paths
    regression = sections[paths.index(["Machine Learning", "Unit 3: Linear Regression"])]
    table = next(block.text for block in regression.blocks if block.kind == "table")
    assert "2" in table and "5" in table and table.count("\n") == 3


def test_formatted_headings_and_paragraphs_stay_whole() -> None:
    notes = (
        b"# Notes\n\n### 3.5 Worked example: A\\* search\n\n"
        b"A **bold** word stays in its sentence.\n"
    )
    sections = split_into_sections(convert_to_docling("text/markdown", notes))
    worked = sections[-1]
    assert worked.heading_path == ["Notes", "3.5 Worked example: A * search"]
    assert [block.text for block in worked.blocks if block.kind == "text"] == [
        "A bold word stays in its sentence."
    ]


def test_word_headings_keep_the_text_nested_under_them() -> None:
    word = Document()
    word.add_heading("Unit 1 Intro", level=1)
    word.add_paragraph("Machine learning learns from data.")
    word.add_heading("1.1 Sub", level=2)
    word.add_paragraph("Supervised learning uses labels.")
    buffer = BytesIO()
    word.save(buffer)
    sections = split_into_sections(convert_to_docling(DOCX, buffer.getvalue()))
    texts = [block.text for section in sections for block in section.blocks]
    assert "Machine learning learns from data." in texts
    assert "Supervised learning uses labels." in texts
    assert sections[-1].heading_path == ["Unit 1 Intro", "1.1 Sub"]
